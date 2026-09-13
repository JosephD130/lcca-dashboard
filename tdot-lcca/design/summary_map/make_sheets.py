"""Generates the three artboards that document what shipped alongside the Summary dashboard:
the setup flow, and the two reference sheets that had never been designed.

Pay_Items and Maintenance Policies are read out of the built workbook, so the artboards cannot
drift from the sheets they show.
"""
import os, re, warnings, html as H
from openpyxl import load_workbook
warnings.filterwarnings('ignore')

HERE = os.path.dirname(os.path.abspath(__file__))
WB = '/home/user/lcca-dashboard/tdot-lcca/TDOA_LCCA_Framework_v1.2.0_ARA_09112026.xlsm'
INK, MUTE, NAVY, BLUE, PALE, GREY = '#1D2733', '#595959', '#1D2733', '#2A78D6', '#EAF2FB', '#D9D9D9'

wb = load_workbook(WB, keep_vba=True)
pi, mp, gi = wb['Pay_Items'], wb['Maintenance Policies'], wb['General Information']
e = lambda v: H.escape(str(v), quote=False) if v is not None else ''

HEAD = '''<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <script src="./support.js"></script>
</head>
<body>
<x-dc>
<helmet>
  <style>
    body { margin: 0; background: #8A8F98; font-family: Arial, Helvetica, sans-serif; color: #1D2733; }
    a { color: #2A78D6; } a:hover { color: #1D5BA6; }
    .sheet { background: #FFFFFF; box-sizing: border-box; padding: 18px 20px 22px; }
    .btn { background: #1D2733; color: #FFFFFF; font-size: 11px; font-weight: bold; padding: 4px 10px;
            display: inline-block; }
    .btn2 { background: #2A78D6; color: #FFFFFF; font-size: 11px; font-weight: bold; padding: 4px 10px;
             display: inline-block; }
    .step { font-size: 10.5px; color: #595959; font-style: italic; }
    .note { font-size: 10px; color: #595959; line-height: 1.45; }
    table.grid { border-collapse: collapse; width: 100%; }
    table.grid td, table.grid th { border: 1px solid #D0D4D9; font-size: 10.5px; padding: 2px 6px;
                                    text-align: left; vertical-align: top; }
    th.hd { background: #1D2733; color: #FFFFFF; font-size: 10px; font-weight: bold; border-color: #1D2733; }
    td.band { background: #EAF2FB; font-weight: bold; }
    td.input { background: #D9D9D9; text-align: right; font-variant-numeric: tabular-nums; }
    td.n { text-align: right; }
    .flag { background: #FFF3CD; color: #7F6000; font-size: 10px; line-height: 1.45; padding: 7px 9px; }
    .title { font-size: 14px; font-weight: bold; letter-spacing: .03em; }
  </style>
</helmet>
<div class="sheet">
'''
TAIL = '''</div>
</x-dc>
</body>
</html>
'''


def write(name, body):
    open(os.path.join(HERE, name), 'w', encoding='utf-8').write(HEAD + body + TAIL)
    print('wrote', name)


# ---------------------------------------------------------------- 1. the setup flow
STEPS = [
    ('1', 'Overview and Instructions', 'Read', 'What the framework does and the rules behind it.', False),
    ('2', 'General Information', 'Fill in', 'The project, the areas and the LCCA parameters. A live line lists whatever is still empty.', True),
    ('3', 'Pay_Items', 'Check', 'The unit costs every alternative is priced from. Edit the grey column if a price is wrong.', True),
    ('4', 'Alternative worksheets', 'Enter', 'Alternative Setup creates one sheet per alternative; type the quantities on each.', True),
    ('5', 'Summary', 'Read', 'Six tiles, the results table, the comparison block and eight charts. It updates by itself.', False),
]
REFS = [('Maintenance Policies', 'What each alternative type does to the pavement and when. The schedules live in the hidden templates.'),
        ('Typical Values', 'The usual range for every input, with the published source beside it.'),
        ('Method', 'Every formula behind the numbers, the rule in plain English, and the assumptions that are a decision.')]

cards = ''
for n, sheet, verb, what, is_input in STEPS:
    cards += (
        '<div style="display: flex; gap: 12px; align-items: flex-start;">'
        '<div style="flex: none; width: 30px; height: 30px; background: %s; color: #FFFFFF; font-size: 15px; '
        'font-weight: bold; display: flex; align-items: center; justify-content: center;">%s</div>'
        '<div style="flex: 1 1 0; min-width: 0; border: 1px solid #D4D8DD; border-left: 3px solid %s; padding: 7px 11px 9px;">'
        '<div style="font-size: 12.5px; font-weight: bold;">%s'
        '<span style="font-weight: normal; font-size: 10px; color: #595959; margin-left: 8px; '
        'text-transform: uppercase; letter-spacing: .05em;">%s</span></div>'
        '<div class="note" style="margin-top: 2px;">%s</div></div></div>'
        % (BLUE if is_input else NAVY, n, BLUE if is_input else NAVY, sheet, verb, what))
    if n != '5':
        cards += ('<div style="margin-left: 14px; width: 2px; height: 12px; background: #C6CBD0;"></div>')

refs = ''.join(
    '<div style="border: 1px solid #D4D8DD; padding: 7px 10px 8px;">'
    '<div style="font-size: 11.5px; font-weight: bold;">%s</div>'
    '<div class="note" style="margin-top: 2px;">%s</div></div>' % (t, d) for t, d in REFS)

write('SetupFlow.dc.html',
      '<div class="title">HOW A PROJECT MOVES THROUGH THE WORKBOOK</div>'
      '<div class="note" style="margin: 5px 0 14px;">The card on General Information lists these five steps and '
      'names the sheet for each. Every sheet in the sequence says which step it is, in the same place, so someone '
      'who lands on one of them mid-way knows where they are. Blue steps are the ones where something is typed.</div>'
      '<div style="display: flex; flex-direction: column;">%s</div>'
      '<div style="margin-top: 18px; font-size: 11px; font-weight: bold; letter-spacing: .05em; '
      'border-bottom: 2px solid #1D2733; padding-bottom: 4px;">REFERENCE, NOT STEPS</div>'
      '<div style="display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 10px; margin-top: 9px;">%s</div>'
      '<div class="note" style="margin-top: 14px;">General Information carries a button for every one of them, so '
      'each is one click from the sheet the user starts on.</div>' % (cards, refs))


# ---------------------------------------------------------------- 2. Pay_Items
BANDS = {4, 26, 34, 37, 44, 48, 54}
rows = ''
for r in range(3, 60):
    if not pi.cell(r, 4).value and r != 3: continue
    part = e(pi.cell(r, 1).value)
    cost = pi.cell(r, 6).value
    if isinstance(cost, str) and re.match(r'^=\d+(\.\d+)?/\d+(\.\d+)?$', cost):
        a, b = cost[1:].split('/'); cost = float(a) / float(b)      # three prices are stored as a division
    txt = ('' if cost in (None, '') else
           (cost if isinstance(cost, str) else '$%0.2f' % cost))
    rows += ('<tr><td%s>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td class="input">%s</td>'
             '<td class="n" style="color:#9AA0A6;">&ndash;</td><td class="n" style="color:#9AA0A6;">&ndash;</td>'
             '<td class="n" style="color:#9AA0A6;">&ndash;</td><td style="color:#595959;">%s</td></tr>'
             % (' class="band"' if r in BANDS else '', part, e(pi.cell(r, 2).value), e(pi.cell(r, 3).value),
                e(pi.cell(r, 4).value), e(pi.cell(r, 5).value), e(txt), e(pi.cell(r, 10).value)))

notes = ''.join('<div class="%s" style="margin-top: 6px;">%s</div>'
                % ('flag' if r == 63 else 'note', e(pi.cell(r, 1).value)) for r in (61, 62, 63, 64))
write('PayItems.dc.html',
      '<div style="display: flex; gap: 8px; align-items: center;">'
      '<span class="btn">&#9668; General Information</span><span class="btn2">Summary</span>'
      '<span class="step" style="margin-left: 8px;">%s</span></div>'
      '<div style="margin-top: 12px; text-align: right; font-size: 10px; font-weight: bold; color: #595959;">'
      'Average Pay Item Unit Cost (Middle / West / East) &mdash; empty in this release</div>'
      '<table class="grid" style="margin-top: 3px;">'
      '<tr><th class="hd">Division</th><th class="hd">Division Type</th><th class="hd">Pay Item No.</th>'
      '<th class="hd" style="width: 34%%;">Pay Item Description</th><th class="hd">Units</th>'
      '<th class="hd">Unit Cost</th><th class="hd">Middle</th><th class="hd">West</th><th class="hd">East</th>'
      '<th class="hd" style="width: 20%%;">Notes</th></tr>%s</table>%s'
      % (e(pi['D1'].value).strip(), rows, notes))


# ---------------------------------------------------------------- 3. Maintenance Policies
tables = ''
for cap in (8, 35, 49, 74):
    end = {8: 33, 35: 47, 49: 72, 74: 86}[cap]
    body = ''
    for r in range(cap + 2, end):
        b, c, d, ee = (mp.cell(r, 2).value, mp.cell(r, 3).value, mp.cell(r, 4).value, mp.cell(r, 5).value)
        if not (b or c): continue
        body += ('<tr><td style="font-weight: bold;">%s</td><td>%s</td><td class="n">%s</td><td class="n">%s</td></tr>'
                 % (e(b), e(c), '' if d is None else '%.4f' % d, e(ee)))
    tables += ('<div style="margin-top: 16px;"><div style="font-size: 12px; font-weight: bold;">%s</div>'
               '<table class="grid" style="margin-top: 4px;">'
               '<tr><th class="hd" style="width: 22%%;">%s</th><th class="hd">Maintenance Item</th>'
               '<th class="hd" style="width: 11%%;">Rate</th><th class="hd" style="width: 11%%;">Year Applied</th></tr>'
               '%s</table></div>'
               % (e(mp.cell(cap, 2).value), e(mp.cell(cap + 1, 2).value), body))

write('MaintenancePolicies.dc.html',
      '<div style="display: flex; gap: 8px; align-items: center;">'
      '<span class="btn">&#9668; General Information</span><span class="btn2">Summary</span></div>'
      '<div style="display: flex; gap: 16px; margin-top: 12px; align-items: flex-start;">'
      '<div style="flex: none; width: 150px; height: 96px; border: 1px dashed #C6CBD0; display: flex; '
      'align-items: center; justify-content: center; font-size: 10px; color: #9AA0A6;">TDOT logo</div>'
      '<div style="flex: 1 1 0; min-width: 0;"><div class="title">%s</div>%s</div></div>%s'
      % (e(mp['C2'].value),
         ''.join('<div class="note" style="margin-top: 5px;">%s</div>' % e(mp['C%d' % r].value)
                 for r in (3, 4, 5, 6)),
         tables))
