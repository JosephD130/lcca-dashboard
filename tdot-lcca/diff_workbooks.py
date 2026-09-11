import openpyxl, json, hashlib, re, sys, zipfile, html
S='/tmp/claude-0/-home-user-lcca-dashboard/444ca227-15d7-51d4-a51e-a0ae3dbe81ed/scratchpad'
U='/root/.claude/uploads/444ca227-15d7-51d4-a51e-a0ae3dbe81ed'
A=U+'/8eaa2547-TDOA_LCCA_Framework_v1.1.2_ARA_Task2_08182026.xlsm'
B=S+'/TDOA_LCCA_Framework_v1.2.0_ARA_09112026.xlsm'
wa=openpyxl.load_workbook(A); wb=openpyxl.load_workbook(B)
def cells(ws):
    d={}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None: d[c.coordinate]=c.value
    return d
diff={}
for name in wb.sheetnames:
    a=cells(wa[name]) if name in wa.sheetnames else {}; b=cells(wb[name])
    ch=[]
    for k in sorted(set(a)|set(b), key=lambda r:(int(re.sub(r'[A-Z]+','',r)), re.sub(r'\d+','',r))):
        va,vb=a.get(k),b.get(k)
        if va!=vb: ch.append({'cell':k,'before':None if va is None else str(va),'after':None if vb is None else str(vb)})
    if ch: diff[name]=ch
tot=sum(len(v) for v in diff.values())
print('changed cells by sheet:',{k:len(v) for k,v in diff.items()},'total',tot)
# data validations
dva=[(str(d.sqref), d.formula1) for d in wa['General Information'].data_validations.dataValidation]
dvb=[(str(d.sqref), d.formula1) for d in wb['General Information'].data_validations.dataValidation]
print('DV before',dva); print('DV after',dvb)
# defined names
print('names before',[n for n in wa.defined_names.keys()] if hasattr(wa.defined_names,'keys') else None)
print('names after',[n for n in wb.defined_names.keys()] if hasattr(wb.defined_names,'keys') else None)
# text box runs in Instructions
def runs(path):
    z=zipfile.ZipFile(path); d=z.read('xl/drawings/drawing3.xml').decode()
    return [html.unescape(t) for t in re.findall(r'<a:t>(.*?)</a:t>',d,re.S)]
ra,rb=runs(A),runs(B)
textdiff=[(i,ra[i],rb[i]) for i in range(len(ra)) if ra[i]!=rb[i]]
print('instruction runs changed:',[i for i,_,_ in textdiff])
hashes={p.split('/')[-1]:hashlib.sha256(open(p,'rb').read()).hexdigest() for p in [A,B,S+'/TDOT_LCCA_Decision_Workbook.xlsx',S+'/Output.bas',U+'/35f08cad-TDOA_LCCA_Framework_v1.1.2_MBT_20260521_LostRevenue_1.xlsm',U+'/527a5482-TDOA_LCCA_Framework_v1.1.2_SRB_20260521_LostRevenue.xlsm']}
json.dump({'diff':diff,'dv':{'before':dva,'after':dvb},'textdiff':textdiff,'hashes':hashes},open(S+'/diff.json','w'),indent=1)
for k,v in diff.items():
    print('==',k); 
    for c in v[:6]: print('  ',c['cell'],'|',(c['before'] or '')[:90],'=>',(c['after'] or '')[:110])
