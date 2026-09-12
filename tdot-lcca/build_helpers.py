"""'Typical Values' reference sheet for the TDOT LCCA framework: what TDOT Aeronautics manages, typical
project geometry, pavement sections, unit costs, maintenance intervals, closure production rates, economic
parameters and the live daily-revenue table. Built with openpyxl and appended to the xlsm by
build_summary.add_plain_sheet. Content rows come from RESEARCH (filled from cited sources)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F_T = Font(name='Arial', size=12, bold=True); F_H = Font(name='Arial', size=10, bold=True); F_B = Font(name='Arial', size=10)
F_N = Font(name='Arial', size=9, italic=True, color='595959'); F_BTN = Font(name='Arial', size=10, bold=True, color='FFFFFF'); F_LINK = Font(name='Arial', size=9, color='1F4E9C', underline='single')
FILL = PatternFill('solid', fgColor='D9D9D9'); FILL_BTN = PatternFill('solid', fgColor='1D2733'); FILL_SEC = PatternFill('solid', fgColor='EAF2FB')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

# Each section: (title, headers, rows, note). Rows are lists of cell values; strings starting with '=' are formulas.
SECTIONS = []   # filled by set_sections() from research; see RESEARCH.md
HINTS = {}      # {General Information row: short hint shown in column F}, filled with SECTIONS

def set_sections(sections): SECTIONS[:] = sections

def build(path):
    wb = Workbook(); ws = wb.active; ws.title = 'Typical Values'
    for col, w in zip('ABCDEFG', [40, 22, 20, 24, 14, 40, 52]): ws.column_dimensions[col].width = w
    c = ws['A1']; c.value = '=HYPERLINK("#\'General Information\'!D9","◄ General Information")'; c.font = F_BTN; c.fill = FILL_BTN; c.alignment = Alignment(horizontal='center', vertical='center'); c.border = BOX
    c = ws['B1']; c.value = '=HYPERLINK("#Summary!G1","Summary")'; c.font = F_BTN; c.fill = PatternFill('solid', fgColor='2A78D6'); c.alignment = Alignment(horizontal='center', vertical='center'); c.border = BOX
    ws.row_dimensions[1].height = 22
    ws['A2'] = 'Reference sheet: nothing here feeds the calculation.'; ws['A2'].font = F_N
    ws['A3'] = 'TYPICAL VALUES AND HELPERS FOR THE INPUTS'; ws['A3'].font = F_T
    ws['A4'] = ('Reference only: nothing on this sheet feeds the calculation. Use it to sanity-check what you type on General Information and the alternative '
                'worksheets. Values are ranges from the cited sources as of the date in the source column; confirm against current bid tabulations for a specific project.'); ws['A4'].font = F_N; ws['A4'].alignment = WRAP
    ws.merge_cells('A4:G4'); ws.row_dimensions[4].height = 30
    ws.page_setup.orientation = 'landscape'
    r = 6
    for title, headers, rows, note in SECTIONS:
        ws.cell(r, 1, title).font = F_H; ws.cell(r, 1).fill = FILL_SEC
        for cc in range(2, 8): ws.cell(r, cc).fill = FILL_SEC
        r += 1
        for i, h in enumerate(headers):
            c = ws.cell(r, 1 + i, h); c.font = F_H; c.fill = FILL; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='center')
        r += 1
        for row in rows:
            for i, v in enumerate(row):
                if isinstance(v, str) and '{row}' in v: v = v.replace('{row}', str(r))
                c = ws.cell(r, 1 + i, v); c.font = F_B; c.border = BOX; c.alignment = WRAP
                if isinstance(v, (int, float)) and not isinstance(v, bool): c.number_format = '0' if (float(v).is_integer() and 1900 < v < 2100 and i == 3) else ('#,##0' if float(v).is_integer() else '#,##0.00')
                elif isinstance(v, str) and v.startswith('=') and i == 1: c.number_format = '#,##0.00'
                if isinstance(v, str) and v.startswith('http'): c.font = F_LINK; c.value = f'=HYPERLINK("{v}","{v[:70]}")' if len(v) < 250 else v
            r += 1
        if note: c = ws.cell(r, 1, note); c.font = F_N; c.alignment = WRAP; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7); ws.row_dimensions[r].height = 26; r += 1
        r += 1
    ws.freeze_panes = 'A5'
    wb.save(path)


def live_sections():
    """Sections computed from the workbook itself (formulas), independent of outside sources."""
    rev_rows = []
    for r in range(2, 19):
        chain = '+'.join(f'SUMIF(RevenueData!$A:$A,A{{row}},RevenueData!${c}:${c})' for c in 'BCDEFGH')
        rev_rows.append([f'=IF(RevenueData!A{r}="","",RevenueData!A{r})', f'=IF(A{{row}}="","",{chain})', f'=IF(A{{row}}="","",B{{row}}*7)', f'=IF(A{{row}}="","",B{{row}}*19)',
                         '=IF(A{row}="","",TRIM(IF(ISNUMBER(RevenueData!B%d),"","100LL ")&IF(ISNUMBER(RevenueData!C%d),"","Jet-A ")&IF(ISNUMBER(RevenueData!G%d),"","tenant ")&IF(ISNUMBER(RevenueData!H%d),"","other")))' % (r, r, r, r), '', 'RevenueData sheet in this workbook (hidden)'])
    return [('Airport daily revenue used for lost-revenue cost (the 17 airports with revenue data; live from RevenueData)',
             ['Airport ID', 'Daily revenue ($/day)', 'Lost revenue, 7-day closure ($)', 'Lost revenue, 19-day closure ($)', 'Columns holding text (counted as $0)', '', 'Source'],
             rev_rows,
             'Same SUMIF chain as General Information D39: text entries in RevenueData (" $-   ", " NA ", descriptions on XNX and M54) count as zero. 7 days is a typical surface-treatment closure and 19 days a mill-and-overlay closure for a 6,000 x 100 ft runway (see production rates).'),
            ('Airports in the General Information dropdown', ['Item', 'Value', '', '', '', '', 'Source'],
             [['Airports selectable in D9', "=COUNTA('General Information'!$L$10:$L$88)", '', '', '', '', 'General Information, hidden columns L:P'],
              ['Airports with revenue data', '=COUNTA(RevenueData!$A$2:$A$60)', '', '', '', '', 'RevenueData'],
              ['Airports with a Middle / West / East region tag', "=COUNTIF('General Information'!$P$10:$P$88,\"Middle\")&\" / \"&COUNTIF('General Information'!$P$10:$P$88,\"West\")&\" / \"&COUNTIF('General Information'!$P$10:$P$88,\"East\")", '', '', '', '', 'General Information, hidden columns L:P']],
             None)]
