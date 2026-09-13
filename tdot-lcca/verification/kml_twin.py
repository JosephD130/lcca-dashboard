"""Python twin of LCCA_KML_Export.bas.

The macro cannot be run in this environment, so the same logic is implemented here against a
populated workbook. It is used to prove the output is well-formed KML with the geometry, the
descriptions and the time stamps in the right places, and to produce a sample file.

usage: python3 kml_twin.py <populated workbook.xlsx> <out.kml>
"""
import sys, math, re, warnings
from openpyxl import load_workbook
warnings.filterwarnings('ignore')

RUNWAY_WIDTH_FT = 100.0
BAR_METRES_PER_MILLION = 180.0
FT_PER_DEG_LAT = 364000.0


def num(v, places): return f'{v:.{places}f}'
def money(v): return '' if not isinstance(v, (int, float)) else '$' + format(round(v), ',')
def X(s): return str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
def X2(s): return str(s).replace(']]>', ']] >')


def corner3(lat, lon, bearing, along_ft, across_ft, alt_m=0.0):
    rad = math.radians(bearing)
    dN = along_ft * math.cos(rad) - across_ft * math.sin(rad)
    dE = along_ft * math.sin(rad) + across_ft * math.cos(rad)
    la = lat + dN / FT_PER_DEG_LAT
    lo = lon + dE / (FT_PER_DEG_LAT * math.cos(math.radians(lat)))
    return f'{num(lo, 6)},{num(la, 6)},{num(alt_m, 1)}'


def runway_bearing(branch):
    m = re.search(r'\d+', str(branch or ''))
    return (int(m.group()) % 36) * 10.0 if m else 0.0


def clean(s):
    o = ''.join(c if (c.isalnum() and c.isascii()) else ('_' if c in ' -_' else '') for c in str(s))
    while '__' in o: o = o.replace('__', '_')
    return o


def build(path_wb, path_out):
    wb = load_workbook(path_wb, data_only=True)
    gi, sm, db = wb['General Information'], wb['Summary'], wb['Database']
    g = lambda ref: gi[ref].value
    lon, lat = sm['AD117'].value, sm['AE117'].value
    assert isinstance(lon, (int, float)) and isinstance(lat, (int, float)), 'no coordinates for this airport'

    L = []
    w = L.append
    w('<?xml version="1.0" encoding="UTF-8"?>')
    w('<kml xmlns="http://www.opengis.net/kml/2.2" xmlns:gx="http://www.google.com/kml/ext/2.2">')
    w('<Document>')
    docname = '%s (%s) - %s' % (g('D9'), g('D10'), g('D16') or g('D22'))
    w('  <name>%s</name>' % X(docname))
    intro = 'Life-cycle cost analysis, %s years at %s%%. Generated from the TDOT Aeronautics LCCA Framework v1.2.0.' % (g('D33'), g('D34'))
    w('  <description>%s</description>' % X(intro))
    for sid, line, poly in [('altLow', 'ff2e8b1f', '662e8b1f'), ('alt', 'ffd6781f', '66d6781f'), ('event', 'ff0e44c1', '660e44c1')]:
        w(f'  <Style id="{sid}">')
        w(f'    <LineStyle><color>{line}</color><width>2</width></LineStyle>')
        w(f'    <PolyStyle><color>{poly}</color></PolyStyle>')
        w(f'    <IconStyle><color>{line}</color><scale>0.9</scale>'
          '</IconStyle>')
        w('  </Style>')
    w('  <Style id="airport"><IconStyle><scale>1.2</scale>'
      '</IconStyle>'
      '<LabelStyle><scale>1.0</scale></LabelStyle></Style>')

    # ---- airport placemark carrying the whole result
    d = f'<h3>{g("D9")} ({g("D10")})</h3><p>{g("D11")}, {g("D13")} region'
    if g('D22'): d += f'<br/>{g("D22")}, {g("D23")}'
    d += f'<br/>Construction {g("D25")}, {g("D33")} years at {g("D34")}%</p>'
    d += ('<table border="1" cellpadding="4" cellspacing="0"><tr><th>Alternative</th><th>Type</th><th>Initial</th>'
          '<th>Net present worth</th><th>Closure days</th><th>Section</th></tr>')
    rows = []
    for r in range(4, 8):
        if sm.cell(r, 7).value:
            rows.append((r, sm.cell(r, 8).value, sm.cell(r, 9).value, sm.cell(r, 10).value,
                         sm.cell(r, 15).value, sm.cell(r, 17).value, sm.cell(79 + r, 16).value, sm.cell(r, 7).value))
    for _, name, typ, init, npw, days, sec, _ws in rows:
        d += (f'<tr><td>{name}</td><td>{typ}</td><td align="right">{money(init)}</td>'
              f'<td align="right">{money(npw)}</td><td align="right">{days}</td><td>{sec}</td></tr>')
    d += '</table>'
    d += f'<p><b>{sm["G9"].value}</b><br/>{sm["G10"].value}</p>'
    w('  <Placemark>')
    w('    <name>%s</name>' % X('%s (%s)' % (g('D9'), g('D10'))))
    w('    <styleUrl>#airport</styleUrl>')
    w(f'    <description><![CDATA[{d}]]></description>')
    w(f'    <Point><coordinates>{num(lon, 6)},{num(lat, 6)},0</coordinates></Point>')
    w('  </Placemark>')

    bearing = runway_bearing(g('D22'))
    length_ft = (g('D26') or 0) * 9.0 / RUNWAY_WIDTH_FT
    yr0, period = int(g('D25')), int(g('D33'))
    best = min(r[4] for r in rows)

    for k, (r, name, typ, init, npw, days, sec, wsname) in enumerate(rows):
        style = 'altLow' if npw == best else 'alt'
        w('  <Folder>')
        w(f'    <name>{X(f"{name} - {money(npw)}" + (" (lowest)" if npw == best else ""))}</name>')
        # schematic runway footprint
        if length_ft > 0:
            off = (k - 1.5) * RUNWAY_WIDTH_FT * 2.2
            hx, hy = length_ft / 2.0, RUNWAY_WIDTH_FT / 2.0
            c = [corner3(lat, lon, bearing, -hx, -hy + off), corner3(lat, lon, bearing, hx, -hy + off),
                 corner3(lat, lon, bearing, hx, hy + off), corner3(lat, lon, bearing, -hx, hy + off)]
            c.append(c[0])
            w('    <Placemark>')
            w(f'      <name>{X(name + " footprint")}</name>')
            w(f'      <styleUrl>#{style}</styleUrl>')
            w(f'      <description><![CDATA[{X2(sec)}<br/>Schematic footprint: {length_ft:,.0f} x {RUNWAY_WIDTH_FT:,.0f} ft '
              'from the area entered, turned to the runway number. Not survey geometry.]]></description>')
            w('      <Polygon><tessellate>1</tessellate><outerBoundaryIs><LinearRing><coordinates>')
            for p in c: w('        ' + p)
            w('      </coordinates></LinearRing></outerBoundaryIs></Polygon>')
            w('    </Placemark>')
        # extruded net present worth bar
        h = (npw / 1e6) * BAR_METRES_PER_MILLION
        s, e = 260.0, 1200.0 + k * 700.0
        c = [corner3(lat, lon, 90.0, e, 0.0, h), corner3(lat, lon, 90.0, e + s, 0.0, h),
             corner3(lat, lon, 90.0, e + s, s, h), corner3(lat, lon, 90.0, e, s, h)]
        c.append(c[0])
        w('    <Placemark>')
        w(f'      <name>{X(f"{name}: {money(npw)}")}</name>')
        w(f'      <styleUrl>#{style}</styleUrl>')
        w('      <Polygon><extrude>1</extrude><altitudeMode>relativeToGround</altitudeMode>')
        w('        <outerBoundaryIs><LinearRing><coordinates>')
        for p in c: w('          ' + p)
        w('      </coordinates></LinearRing></outerBoundaryIs></Polygon>')
        w('    </Placemark>')
        # timed events
        ws = wb[wsname]
        hdr = npw_row = 0
        for rr in range(30, 61):
            v = str(ws.cell(rr, 2).value or '').strip()
            if v == 'Item': hdr = rr
            if v == 'Net Present Worth': npw_row = rr
        daily = ws['F2'].value if isinstance(ws['F2'].value, (int, float)) else 0
        n = 0
        for rr in range(hdr + 1, npw_row):
            item = str(ws.cell(rr, 2).value or '').strip()
            if not item or 'indirect' in item.lower() or item == 'Item': continue
            yv, cv = ws.cell(rr, 3).value, ws.cell(rr, 4).value
            if not isinstance(yv, (int, float)) or not isinstance(cv, (int, float)): continue
            yoff = 0 if yv >= 1000 else yv
            if cv == 0 or yoff > period: continue
            days_ = 0
            nxt = ws.cell(rr + 1, 4).value
            if daily and isinstance(nxt, (int, float)) and 'indirect' in str(ws.cell(rr + 1, 2).value or '').lower():
                days_ = nxt / daily
            pos = (n / 8.0 - 0.45) * length_ft
            year = yr0 + int(yoff)
            w('    <Placemark>')
            w(f'      <name>{X(f"{item} {year}")}</name>')
            w('      <styleUrl>#event</styleUrl>')
            w(f'      <TimeSpan><begin>{year}-01-01</begin><end>{year}-12-31</end></TimeSpan>')
            w(f'      <description><![CDATA[<b>{X2(name)}</b><br/>{X2(item)}<br/>Year {year} ({int(yoff)} years after '
              f'construction)<br/>Cost in the year: {money(cv)}' +
              (f'<br/>Runway closed about {days_:.0f} days' if days_ > 0 else '') + ']]></description>')
            w(f'      <Point><coordinates>{corner3(lat, lon, bearing, pos, (k - 1.5) * 220.0, 0.0)}</coordinates></Point>')
            w('    </Placemark>')
            n += 1
        w('  </Folder>')
    w('</Document>')
    w('</kml>')
    open(path_out, 'w', encoding='utf-8').write('\n'.join(L) + '\n')
    return path_out, clean(f'{g("D10")}_{g("D22")}_LCCA') + '.kml'


if __name__ == '__main__':
    out, name = build(sys.argv[1], sys.argv[2])
    print('wrote', out, '(the macro would name it', name + ')')
