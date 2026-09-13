"""Static verification of the delivered workbook: package hygiene, sheet inventory, Summary layout,
charts (axis titles, legends, hidden-cell plotting), the Method and Typical Values sheets, and the
pavement-section block. Reads the file as a zip and with openpyxl; no Excel or LibreOffice needed.

usage: python3 verify_build.py <workbook.xlsm>
"""
import sys, re, zipfile, warnings
from openpyxl import load_workbook
warnings.filterwarnings('ignore')

WB = sys.argv[1] if len(sys.argv) > 1 else 'TDOA_LCCA_Framework_v1.2.0_ARA_09112026.xlsm'
z = zipfile.ZipFile(WB)
names = z.namelist()
rd = lambda n: z.read(n).decode('utf-8', 'replace')
wb = load_workbook(WB, keep_vba=True)
ws = wb['Summary']
fails = []


def check(name, ok, detail=''):
    print(('PASS  ' if ok else 'FAIL  ') + name + (('   |   ' + str(detail)[:160]) if detail else ''))
    if not ok: fails.append(name)


print('=' * 78); print('PACKAGE'); print('=' * 78)
check('zip test: every part parses', all(z.testzip() is None for _ in [0]))
check('VBA project still present', 'xl/vbaProject.bin' in names)
check('all 106 ActiveX controls kept', sum(1 for n in names if re.match(r'xl/activeX/activeX\d+\.xml$', n)) == 106,
      sum(1 for n in names if re.match(r'xl/activeX/activeX\d+\.xml$', n)))
check('every ActiveX binary keeps a content type', rd('[Content_Types].xml').count('/xl/activeX/activeX') == 106 * 2 or
      len(re.findall(r'PartName="/xl/activeX/activeX\d+\.bin"', rd('[Content_Types].xml'))) == 106)
check('printer settings removed', not [n for n in names if 'printerSettings' in n])
check('no page setup still points at a printer part', not re.search(r'<pageSetup[^>]*r:id=', ''.join(rd(n) for n in names if n.startswith('xl/worksheets/sheet'))))
check('Power Query stub removed', not [n for n in names if n.startswith('customXml/')])
check('cached SharePoint path removed', 'absPath' not in rd('xl/workbook.xml'))
core = rd('docProps/core.xml')
check('document properties carry no personal names', 'Jim Bruinsma' not in core and 'Marti Denis' not in core,
      re.findall(r'<dc:creator>([^<]*)', core) + re.findall(r'<cp:lastModifiedBy>([^<]*)', core))
check('external links removed', not [n for n in names if n.startswith('xl/externalLinks')] and 'externalReferences' not in rd('xl/workbook.xml'))
check('calculation chain dropped (full recalc on open)', 'xl/calcChain.xml' not in names and 'fullCalcOnLoad="1"' in rd('xl/workbook.xml'))
check('no cell carries an invalid style id', 's="None"' not in ''.join(rd(n) for n in names if n.startswith('xl/worksheets/sheet')))

print(); print('=' * 78); print('SHEETS'); print('=' * 78)
expect = ['TMP(NewPCC)_IndirectCost', 'Overview', 'Instructions', 'General Information', 'Pay_Items', 'Database',
          'Maintenance Policies', 'TMP(NewHMA)_IndirectCost', 'RevenueData', 'TMP(NewHMA)', 'TMP(NewPCC)',
          'TMP(HMARehab)', 'TMP(PCCRehab)', 'Summary', 'Typical Values', 'Method']
check('sheet order unchanged, Typical Values and Method appended last', wb.sheetnames == expect, wb.sheetnames)
check('the four templates stay hidden', all(wb[s].sheet_state != 'visible' for s in expect if s.startswith('TMP(')))
check('Database and RevenueData stay hidden', wb['Database'].sheet_state == 'hidden' and wb['RevenueData'].sheet_state == 'hidden')

print(); print('=' * 78); print('SUMMARY LAYOUT'); print('=' * 78)
check('columns A to F hidden (the block the form writes)', all(ws.column_dimensions[c].hidden for c in 'ABCDEF'),
      {c: ws.column_dimensions[c].hidden for c in 'ABCDEFG'})
check('column G onward visible', not ws.column_dimensions['G'].hidden)
check('navigation buttons sit in the first visible column', 'HYPERLINK' in str(ws['G1'].value) and 'HYPERLINK' in str(ws['H1'].value),
      (ws['G1'].value, ws['H1'].value))
check('title moved beside the buttons', ws['J1'].value == 'LCCA SUMMARY')
check('note explains the hidden block', 'hidden' in str(ws['G2'].value) and 'Unhide' in str(ws['G2'].value))
check('results table header intact', [ws.cell(3, c).value for c in range(7, 19)] ==
      ['Worksheet', 'Alternative', 'Type', 'Initial construction', 'Maintenance PW', 'Rehabilitation PW', 'Lost revenue PW',
       'Salvage PW', 'Net present worth', 'vs. lowest NPW', 'Closure days in period', 'Runway availability'],
      [ws.cell(3, c).value for c in range(7, 19)])
check('results table reads the Database sheet by INDEX (survives a row deletion)',
      all(str(ws.cell(r, 7).value).startswith('=IF(INDEX(Database!$D:$D') for r in range(4, 8)))
check('comparison block present (RealCost layout)', str(ws['G12'].value).startswith('COMPARISON') and ws['G13'].value == 'Alternative')
check('section block present with both description strings', ws['P82'].value == 'Section from the quantities'
      and ws['Q82'].value == 'Same quantities over mainline + shoulder')
check('asphalt unit weight is an input cell', ws['J81'].value == 145)
check('chart-data block labelled do not edit', 'do not edit' in str(ws.cell(1, 23).value))
check('an alternative that does not exist gets no name', all('""' in str(ws.cell(4, c).value) for c in range(24, 28)),
      ws.cell(4, 24).value)
wbx = rd('xl/workbook.xml')
check('print areas defined for the Summary and General Information',
      'Summary!$A$1:$V$112' in wbx and "'General Information'!$A$1:$J$48" in wbx,
      re.findall(r'<definedName name="_xlnm.Print_Area"[^>]*>([^<]*)', wbx))

print(); print('=' * 78); print('CHARTS'); print('=' * 78)
charts = sorted((n for n in names if re.match(r'xl/charts/chart\d+\.xml$', n)), key=lambda n: int(re.search(r'\d+', n.split('/')[-1]).group()))
check('fourteen charts: six original, seven new on the Summary, one locator map', len(charts) == 14, len(charts))
for n in charts:
    c = rd(n)
    ts = re.findall(r'<a:t>([^<]*)</a:t>', c)
    has_x = any(t in ('Calendar year', 'Alternative', 'Discount rate (%)', 'Scenario', 'Longitude') for t in ts)
    has_y = any(t in ('Cost ($)', 'Present worth ($)', 'Net present worth ($)', 'Spend, undiscounted ($)',
                      'Cumulative discounted cost ($)', 'Closure days', 'Thickness (inches)', 'Latitude') for t in ts)
    check('%s names both axes' % n.split('/')[-1], has_x and has_y, ts[:4])
summary_charts = charts[6:13]      # the seven cost and section charts; the locator map carries no legend
check('the seven Summary charts put the legend beside the plot',
      all(re.search(r'<legendPos val="r"/>', rd(n)) for n in summary_charts),
      [re.findall(r'<legendPos val="(\w)"/>', rd(n)) for n in summary_charts])
check('money axes read to one decimal in millions',
      sum('$#,##0.0,,&quot;M&quot;' in rd(n) or '$#,##0.0,,"M"' in rd(n) for n in charts) >= 5)
c6 = rd([n for n in charts if n.endswith('chart6.xml')][0])
check('the original Alternatives Comparison chart plots the hidden columns', 'plotVisOnly val="0"' in c6)
check('the alternative-sheet charts keep calendar years on the category axis',
      all('Calendar year' in rd(n) for n in charts[:5]))

print(); print('=' * 78); print('LOCATOR MAP AND PROJECT IDENTITY'); print('=' * 78)
check('project identity line at the top of the Summary', 'D$9' in str(ws['L1'].value) and 'construction' in str(ws['L1'].value))
check('map data block labelled and outside the print area', 'MAP DATA' in str(ws.cell(115, 23).value))
coords = [(ws.cell(r, 25).value, ws.cell(r, 26).value) for r in range(117, 196)]
check('79 airports in the map block, 74 with published coordinates',
      len(coords) == 79 and sum(1 for a, b in coords if isinstance(a, (int, float)) and isinstance(b, (int, float))) == 74,
      (len(coords), sum(1 for a, b in coords if isinstance(a, (int, float)))))
check('coordinates fall inside Tennessee',
      all(-90.5 < a < -81.5 and 34.9 < b < 36.8 for a, b in coords if isinstance(a, (int, float))))
border = [(ws.cell(r, 34).value, ws.cell(r, 35).value) for r in range(117, 330)]
bpts = [(a, b) for a, b in border if isinstance(a, (int, float))]
check('state outline embedded with gaps between segments', len(bpts) == 192 and len(border) > len(bpts), (len(bpts), len(border)))
check('the selected airport is looked up, not typed',
      str(ws.cell(117, 30).value).startswith('=IFERROR(IF(INDEX(') and str(ws.cell(117, 31).value).startswith('=IFERROR(IF(INDEX('))
mapch = [c for c in wb['Summary']._charts if c.tagname == 'scatterChart']
check('locator map has the outline, the three Grand Divisions and this project', len(mapch) == 1 and len(mapch[0].series) == 5,
      [(c.tagname, len(c.series)) for c in wb['Summary']._charts])
check('airports grouped by division so each is its own series',
      [ws.cell(r, 27).value for r in (117, 137, 160, 190)] == ['West', 'West', 'Middle', 'East'],
      [ws.cell(r, 27).value for r in (117, 137, 160, 190)])
check('legend in cells, one per division plus this project',
      [ws.cell(17, c).value for c in range(19, 23)] == ['\u25a0 West', '\u25a0 Middle', '\u25a0 East', '\u25a0 This project'],
      [ws.cell(17, c).value for c in range(19, 23)])
check('card header band across S2:V2', ws['S2'].value == 'PROJECT LOCATION' and ws['S2'].fill.fgColor.rgb.endswith('1D2733'))
check('pricing basis stated beside the map', 'Unit Cost' in str(ws['S28'].value) and 'empty' in str(ws['S28'].value))
check('map axes are named and their degree labels suppressed',
      mapch and mapch[0].x_axis.numFmt.formatCode == ';;;' and mapch[0].y_axis.numFmt.formatCode == ';;;')
check('project block names airport, county, region, coordinates, elevation',
      [ws.cell(20 + k, 19).value for k in range(7)] ==
      ['Airport', 'City / county', 'TDOT region', 'Coordinates', 'Elevation', 'Branch / project', 'Mainline area'],
      [ws.cell(20 + k, 19).value for k in range(7)])
check('the sheet says where the coordinates come from and how to export KML',
      'embedded' in str(ws['S29'].value) and 'ExportLCCAKML' in str(ws['S30'].value))
check('runway width for the export footprint is an input cell', ws['S27'].value == 'Runway width, ft'
      and ws['T27'].value == 100, (ws['S27'].value, ws['T27'].value))
check('the three notes beside the map are merged so they do not run into the chart data',
      all(str(rng) in [str(m) for m in ws.merged_cells.ranges] for rng in ('S28:V28', 'S29:V29', 'S30:V30')),
      [str(m) for m in ws.merged_cells.ranges][-4:])
check('print area widened to take the map', 'Summary!$A$1:$V$112' in rd('xl/workbook.xml'))

print(); print('=' * 78); print('METHOD AND TYPICAL VALUES'); print('=' * 78)
me = wb['Method']
check('Method sheet has its navigation buttons', 'HYPERLINK' in str(me['A1'].value) and 'HYPERLINK' in str(me['B1'].value))
check('Method sheet states it feeds nothing', 'nothing here feeds' in str(me['A2'].value).lower())
titles = [me.cell(r, 1).value for r in range(1, me.max_row + 1) if isinstance(me.cell(r, 1).value, str) and re.match(r'^\d+\. ', me.cell(r, 1).value)]
check('eleven sections, in order', len(titles) == 11 and titles[0].startswith('1.') and titles[-1].startswith('11.'), titles)
formula_cells = [me.cell(r, 3).value for r in range(1, me.max_row + 1) if isinstance(me.cell(r, 3).value, str) and me.cell(r, 3).value.startswith('=')]
live = [c for c in formula_cells if "'General Information'!$D$3" in c or 'COUNTA(Database' in c]
check('the live block reads the workbook (period, rate, CRF, discount factor, alternatives)', len(live) >= 5, len(live))
shown = [me.cell(r, 3) for r in range(1, me.max_row + 1) if isinstance(me.cell(r, 3).value, str)
         and me.cell(r, 3).value.startswith('=') and me.cell(r, 3).data_type == 's']
check('documented formulas are stored as text, not evaluated', len(shown) >= 20, len(shown))
check('assumptions section names the four open questions',
      all(k in ''.join(str(me.cell(r, 1).value) for r in range(1, me.max_row + 1)) for k in
          ['Engineering on initial construction', 'Salvage is not symmetric', 'Lost revenue is gross revenue', 'PCC joint length']))
tv = wb['Typical Values']
srcs = [tv.cell(r, 7).value for r in range(1, tv.max_row + 1) if isinstance(tv.cell(r, 7).value, str) and tv.cell(r, 7).value.startswith('http')]
check('Typical Values sources are plain text, not hyperlinks', len(srcs) >= 20 and not any('HYPERLINK' in str(tv.cell(r, 7).value)
      for r in range(1, tv.max_row + 1)), len(srcs))
hl = sum(len(sh._hyperlinks) for sh in wb.worksheets)
check('only the original tn.gov hyperlink object remains', hl <= 1, hl)

print(); print('=' * 78); print('GENERAL INFORMATION'); print('=' * 78)
gi = wb['General Information']
check('analysis period default 30 years', gi['D33'].value == 30)
check('discount rate default 3 percent', gi['D34'].value == 3)
gix = rd('xl/worksheets/sheet4.xml')
check('airport dropdown covers all 79 airports (D9 list = L10:L88)', '<formula1>$L$10:$L$88</formula1>' in gix,
      re.findall(r'<formula1>([^<]*)</formula1>', gix)[:4])
check('79 airports and 17 with revenue data behind it',
      sum(1 for r in range(10, 89) if gi.cell(r, 12).value) == 79 and sum(1 for r in range(2, 60) if wb['RevenueData'].cell(r, 1).value) == 17,
      (sum(1 for r in range(10, 89) if gi.cell(r, 12).value), sum(1 for r in range(2, 60) if wb['RevenueData'].cell(r, 1).value)))
check('three navigation buttons: Summary, Typical Values, Method',
      all('HYPERLINK' in str(gi[c].value) for c in ('D44', 'D46', 'D48')), [gi[c].value for c in ('D44', 'D46', 'D48')])
check('how-to card and live status line present', str(gi['F2'].value or '').startswith('HOW TO USE') and 'Still needed' in str(gi['F9'].value))
check('input hints beside the parameters', sum(1 for r in (25, 26, 27, 28, 33, 34, 36, 37, 38) if gi.cell(r, 6).value) == 9)

print(); print('=' * 78); print('SPELLING IN THE ISSUED TEXT'); print('=' * 78)
ov = rd('xl/drawings/drawing2.xml'); ss = rd('xl/sharedStrings.xml')
for bad, good, part in [('Adminimstration', 'Administration', ov), ('clossures', 'closures', ov),
                        ('associeted', 'associated', ov), ('Intial Construction', 'Initial Construction', ss)]:
    check('%s corrected to %s' % (bad, good), bad not in part and good in part)


print(); print('=' * 78)
print('%d checks, %d failed' % (len(fails) + sum(1 for _ in []) + 0 if False else 0, len(fails)) if False else
      ('ALL PASS' if not fails else 'FAILED: ' + '; '.join(fails)))
sys.exit(1 if fails else 0)
